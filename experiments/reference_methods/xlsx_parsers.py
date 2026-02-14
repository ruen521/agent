from __future__ import annotations

from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any


@dataclass
class SalesDailyRow:
    as_of_date: date
    asin: str
    units: float
    orders_count: float
    amount: float


@dataclass
class InventorySnapshotAsinRow:
    snapshot_date: date
    asin: str
    total_stock: float
    sellable_stock: float
    inbound_stock: float
    avg_unit_cost: float
    total_inventory_cost: float
    total_goods_value: float
    category: str
    product_name: str
    store_count: int
    warehouse_count: int
    raw_row_count: int


_REQUIRED_SALES_SHEETS = ("销量", "订单量", "销售额")


def _to_float(value: Any) -> float:
    if value is None or value == "":
        return 0.0
    try:
        return float(value)
    except Exception:
        return 0.0


def _parse_header_dates(headers: list[Any]) -> list[tuple[int, date]]:
    cols: list[tuple[int, date]] = []
    for idx, header in enumerate(headers):
        if isinstance(header, str):
            try:
                cols.append((idx, datetime.fromisoformat(header).date()))
            except Exception:
                continue
    cols.sort(key=lambda item: item[1])
    return cols


def parse_sales_workbook(path: str | Path) -> list[SalesDailyRow]:
    try:
        from openpyxl import load_workbook
    except Exception as exc:  # pragma: no cover
        raise RuntimeError("openpyxl is required: pip install openpyxl") from exc

    workbook_path = Path(path)
    if not workbook_path.exists():
        raise FileNotFoundError(f"Workbook not found: {workbook_path}")

    wb = load_workbook(workbook_path, data_only=True, read_only=True)
    for sheet in _REQUIRED_SALES_SHEETS:
        if sheet not in wb.sheetnames:
            raise ValueError(f"Missing sheet in sales workbook: {sheet}")

    sheet_data: dict[str, dict[str, dict[date, float]]] = {}
    canonical_dates: list[date] | None = None

    for sheet in _REQUIRED_SALES_SHEETS:
        ws = wb[sheet]
        headers = list(next(ws.iter_rows(min_row=1, max_row=1, values_only=True)))
        date_columns = _parse_header_dates(headers)
        if not date_columns:
            raise ValueError(f"No date columns found in sheet: {sheet}")
        dates = [d for _, d in date_columns]
        if canonical_dates is None:
            canonical_dates = dates
        elif canonical_dates != dates:
            raise ValueError(f"Date columns mismatch across sheets: {sheet}")

        asin_map: dict[str, dict[date, float]] = {}
        for row in ws.iter_rows(min_row=2, values_only=True):
            asin = str(row[0] or "").strip() if len(row) > 0 else ""
            if not asin:
                continue
            values: dict[date, float] = {}
            for idx, day in date_columns:
                raw = row[idx] if idx < len(row) else 0
                values[day] = _to_float(raw)
            asin_map[asin] = values
        sheet_data[sheet] = asin_map

    assert canonical_dates is not None

    all_asins = set(sheet_data["销量"].keys()) | set(sheet_data["订单量"].keys()) | set(sheet_data["销售额"].keys())
    rows: list[SalesDailyRow] = []
    for asin in sorted(all_asins):
        units_map = sheet_data["销量"].get(asin, {})
        orders_map = sheet_data["订单量"].get(asin, {})
        amount_map = sheet_data["销售额"].get(asin, {})
        for day in canonical_dates:
            rows.append(
                SalesDailyRow(
                    as_of_date=day,
                    asin=asin,
                    units=_to_float(units_map.get(day, 0.0)),
                    orders_count=_to_float(orders_map.get(day, 0.0)),
                    amount=_to_float(amount_map.get(day, 0.0)),
                )
            )
    return rows


def _parse_snapshot_date(raw: Any) -> date:
    if isinstance(raw, datetime):
        return raw.date()
    if isinstance(raw, date):
        return raw
    if isinstance(raw, str):
        text = raw.strip()
        if not text:
            raise ValueError("empty snapshot datetime")
        try:
            return datetime.fromisoformat(text).date()
        except Exception as exc:
            raise ValueError(f"invalid snapshot datetime: {text}") from exc
    raise ValueError(f"unsupported snapshot datetime type: {type(raw)}")


def _mode_non_empty(values: list[str]) -> str:
    cleaned = [value.strip() for value in values if value and value.strip()]
    if not cleaned:
        return ""
    counter = Counter(cleaned)
    return counter.most_common(1)[0][0]


def _first_non_empty(values: list[str]) -> str:
    for value in values:
        if value and value.strip():
            return value.strip()
    return ""


def parse_fba_inventory_workbook(path: str | Path) -> list[InventorySnapshotAsinRow]:
    try:
        from openpyxl import load_workbook
    except Exception as exc:  # pragma: no cover
        raise RuntimeError("openpyxl is required: pip install openpyxl") from exc

    workbook_path = Path(path)
    if not workbook_path.exists():
        raise FileNotFoundError(f"Workbook not found: {workbook_path}")

    wb = load_workbook(workbook_path, data_only=True, read_only=True)
    if not wb.sheetnames:
        raise ValueError("FBA inventory workbook has no sheets")

    ws = wb[wb.sheetnames[0]]
    headers = list(next(ws.iter_rows(min_row=1, max_row=1, values_only=True)))
    index = {str(name).strip(): idx for idx, name in enumerate(headers) if name is not None}
    required_cols = [
        "ASIN",
        "仓库",
        "店铺",
        "品名",
        "分类",
        "总库存",
        "可售",
        "在途",
        "平均采购成本",
        "库存成本",
        "货值",
        "更新时间",
    ]
    missing = [col for col in required_cols if col not in index]
    if missing:
        raise ValueError(f"Missing columns in FBA inventory workbook: {','.join(missing)}")

    grouped: dict[tuple[date, str], dict[str, Any]] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        asin = str(row[index["ASIN"]] or "").strip()
        if not asin:
            continue
        snapshot_date = _parse_snapshot_date(row[index["更新时间"]])
        key = (snapshot_date, asin)
        bucket = grouped.setdefault(
            key,
            {
                "total_stock": 0.0,
                "sellable_stock": 0.0,
                "inbound_stock": 0.0,
                "avg_unit_cost_sum": 0.0,
                "avg_unit_cost_count": 0,
                "total_inventory_cost": 0.0,
                "total_goods_value": 0.0,
                "stores": set(),
                "warehouses": set(),
                "categories": [],
                "product_names": [],
                "raw_row_count": 0,
            },
        )

        bucket["total_stock"] += _to_float(row[index["总库存"]])
        bucket["sellable_stock"] += _to_float(row[index["可售"]])
        bucket["inbound_stock"] += _to_float(row[index["在途"]])

        unit_cost = _to_float(row[index["平均采购成本"]])
        if unit_cost != 0:
            bucket["avg_unit_cost_sum"] += unit_cost
            bucket["avg_unit_cost_count"] += 1

        bucket["total_inventory_cost"] += _to_float(row[index["库存成本"]])
        bucket["total_goods_value"] += _to_float(row[index["货值"]])
        bucket["stores"].add(str(row[index["店铺"]] or "").strip())
        bucket["warehouses"].add(str(row[index["仓库"]] or "").strip())
        bucket["categories"].append(str(row[index["分类"]] or "").strip())
        bucket["product_names"].append(str(row[index["品名"]] or "").strip())
        bucket["raw_row_count"] += 1

    rows: list[InventorySnapshotAsinRow] = []
    for (snapshot_date, asin), bucket in sorted(grouped.items(), key=lambda item: (item[0][0], item[0][1])):
        count = int(bucket["avg_unit_cost_count"])
        avg_unit_cost = 0.0 if count == 0 else float(bucket["avg_unit_cost_sum"]) / count
        rows.append(
            InventorySnapshotAsinRow(
                snapshot_date=snapshot_date,
                asin=asin,
                total_stock=round(float(bucket["total_stock"]), 4),
                sellable_stock=round(float(bucket["sellable_stock"]), 4),
                inbound_stock=round(float(bucket["inbound_stock"]), 4),
                avg_unit_cost=round(avg_unit_cost, 4),
                total_inventory_cost=round(float(bucket["total_inventory_cost"]), 4),
                total_goods_value=round(float(bucket["total_goods_value"]), 4),
                category=_mode_non_empty(list(bucket["categories"])),
                product_name=_first_non_empty(list(bucket["product_names"])),
                store_count=len([x for x in bucket["stores"] if x]),
                warehouse_count=len([x for x in bucket["warehouses"] if x]),
                raw_row_count=int(bucket["raw_row_count"]),
            )
        )
    return rows
